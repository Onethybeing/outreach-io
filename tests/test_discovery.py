import io
import uuid
from datetime import datetime, timedelta, timezone

import docx
import httpx
import pypdfium2
import pytest
from sqlalchemy import select

from app import llm, prompts, telemetry, vault
from app.discovery import extract, graph, tavily
from app.discovery.normalize import company_key, linkedin_profile_url, title_names_other_company, website_domain
from app.models import (
    Candidate, Contact, Resume, Run, RunEvent, RunStatus, Startup, UserRole,
)

PROFILE = {
    "name": "Priya", "current_role": "ML Engineer", "seniority": "mid", "years_experience": "4",
    "skills": ["Python", None, "RAG"], "domains": ["AI infra"], "target_titles": ["ML Engineer"],
}


@pytest.fixture(autouse=True)
def no_tracing(monkeypatch):
    monkeypatch.setattr(telemetry, "langfuse_client", lambda db: None)


@pytest.fixture
def resume(db):
    r = Resume(filename="priya.pdf", storage_path="storage/does-not-matter.pdf",
               extracted_text="Priya — ML Engineer. " * 40, extraction_method="pdf_text")
    db.add(r)
    db.commit()
    return r


def _make_run(db, resume, num_startups=3, num_kdms=2, status=RunStatus.pending):
    prompts.seed_defaults(db)
    active = prompts.active_versions(db)
    run = Run(resume_id=resume.id, num_startups=num_startups, num_kdms_per_company=num_kdms, status=status,
              prompt_versions={n: str(active[n].id) for n in graph.PROMPT_NODES})
    db.add(run)
    db.commit()
    return run


class FakeWorld:
    """Scripted search + LLM answers, including the messy cases the graph must clean up."""

    def __init__(self):
        self.llm_calls: list[str] = []
        self.searches: list[str] = []
        self.startups = [
            {"name": "Acme Vector", "website": "https://www.acmevector.example", "domain": "AI infra",
             "description": "Retrieval infra", "relevance": 0.9, "source_url": "https://news.example/acme"},
            {"name": "Acme Vector, Inc.", "website": "acmevector.example", "relevance": 0.8},  # same company
            {"name": "Old Co", "website": "oldco.example", "relevance": 0.85},  # found in an earlier run
            {"name": "Beta Labs", "website": "betalabs.example", "relevance": 0.7, "source_url": "https://made-up.example"},
            {"name": "", "relevance": 1.0},
        ]
        self.fail_on: str | None = None
        self.fail_kdms_on_call: int | None = None
        self.kdm_calls = 0
        self.research_calls = 0
        self.kdms: dict | None = None  # set to override what find_kdms returns
        self.brief_override: dict | None = None
        self.fail_searches: tuple[str, ...] = ()  # substrings of queries that should fail

    def search(self, db, query, max_results=5, include_domains=None, content_chars=400):
        self.searches.append(query)
        if not include_domains:
            return [{"title": "Acme raises", "url": "https://news.example/acme", "content": "..."},
                    {"title": "Beta", "url": "https://news.example/beta", "content": "..."}]
        if "Acme" in query:
            return [{"title": "Other - CEO at Acme Vectorless", "url": "https://www.linkedin.com/in/lookalike", "content": ""},
                    {"title": "Jane - CTO - Acme", "url": "https://uk.linkedin.com/in/Jane-Doe?trk=x", "content": ""},
                    {"title": "Sam - CEO - Acme", "url": "https://www.linkedin.com/in/sam-founder", "content": ""},
                    {"title": "Existing - Acme", "url": "https://www.linkedin.com/in/already-contact", "content": ""}]
        return [{"title": "Sam - advisor - Beta", "url": "https://www.linkedin.com/in/sam-founder", "content": ""},
                {"title": "Lee - Head of Talent - Beta", "url": "https://www.linkedin.com/in/lee-talent", "content": ""}]

    def search_many(self, db, queries, max_results=5, include_domains=None, content_chars=400):
        """Same answers as `search`, in the batched shape the graph now calls. A query listed in
        `fail_searches` comes back as an error in its own slot, as the real one does."""
        out = []
        for query in queries:
            if any(marker in query for marker in self.fail_searches):
                out.append(tavily.SearchError("Tavily usage limit reached (HTTP 432)"))
            else:
                out.append(self.search(db, query, max_results, include_domains, content_chars))
        return out

    def complete_json(self, db, model, prompt, temperature, max_tokens=4096, **kwargs):
        name = kwargs["name"]
        self.llm_calls.append(name)
        if name == self.fail_on:
            raise llm.LLMError("Groq error: model overloaded")
        if name == "parse_resume":
            return PROFILE
        if name == "build_search_queries":
            return {"queries": ["AI infra startups", "ai infra startups ", "RAG startups hiring", 7]}
        if name == "discover_startups":
            return {"startups": self.startups}
        if name == "research_startup":
            self.research_calls += 1
            if self.brief_override is not None:
                return self.brief_override
            return {
                "what_they_do": "Builds retrieval infrastructure.",
                "product": "A hosted vector search API",
                "stage": "Series A",
                "team_size": 18,
                # The second item cites a URL the search never returned: it must be dropped.
                "recent_news": [
                    {"fact": "Raised a Series A", "url": "https://news.example/acme"},
                    {"fact": "Opened an office on the moon", "url": "https://invented.example/nope"},
                ],
                "tech_signals": ["RAG", "vector search"],
                "hiring_signals": "ML engineers",
                "overlap_with_candidate": ["Has built RAG pipelines"],
            }
        if name == "find_kdms":
            self.kdm_calls += 1
            if self.kdm_calls == self.fail_kdms_on_call:
                raise llm.LLMError("Groq error: failed on a later startup")
            if self.kdms is not None:
                return self.kdms
            if "Acme" in prompt:
                return {"people": [
                    {"name": "Other", "title": "CEO at Acme Vectorless", "linkedin_url": "https://www.linkedin.com/in/lookalike"},
                    {"name": "Jane Doe", "title": "CTO", "linkedin_url": "https://www.linkedin.com/in/jane-doe/", "why": "tech lead"},
                    {"name": "Ghost", "title": "CEO", "linkedin_url": "https://www.linkedin.com/in/not-in-search"},
                    {"name": "Already", "title": "Founder", "linkedin_url": "https://linkedin.com/in/already-contact"},
                    {"name": "Sam", "title": "CEO", "linkedin_url": "https://www.linkedin.com/in/sam-founder"},
                ]}
            return {"people": [
                {"name": "Sam", "title": "Advisor", "linkedin_url": "https://www.linkedin.com/in/sam-founder"},
                {"name": "Lee", "title": "Head of Talent", "linkedin_url": "https://www.linkedin.com/in/lee-talent"},
            ]}
        raise AssertionError(name)


@pytest.fixture
def world(monkeypatch):
    w = FakeWorld()
    monkeypatch.setattr(tavily, "search", w.search)
    monkeypatch.setattr(tavily, "search_many", w.search_many)
    monkeypatch.setattr(llm, "complete_json", w.complete_json)
    return w


def _events(db, run):
    return list(db.scalars(select(RunEvent).where(RunEvent.run_id == run.id).order_by(RunEvent.id)))


def _earlier_run_with_startup_and_contact(db, resume):
    earlier = _make_run(db, resume, status=RunStatus.completed)
    old = Startup(run_id=earlier.id, name="Old Co", website="https://oldco.example")
    db.add(old)
    db.flush()
    db.add(Candidate(run_id=earlier.id, startup_id=old.id, name="Proposed Before",
                     linkedin_url="https://www.linkedin.com/in/proposed-before"))
    db.add(Contact(startup_id=old.id, run_id=earlier.id, resume_id=resume.id, cv_used_id=resume.id,
                   name="Already", linkedin_url="https://www.linkedin.com/in/already-contact"))
    db.commit()


def test_happy_path_cleans_and_dedupes(db, resume, world):
    _earlier_run_with_startup_and_contact(db, resume)
    run = _make_run(db, resume, num_startups=3, num_kdms=2)

    graph.execute_run(db, run.id)
    db.refresh(run)

    assert run.status == RunStatus.completed, run.error
    assert run.started_at and run.finished_at and run.usage is not None

    startups = list(db.scalars(select(Startup).where(Startup.run_id == run.id)))
    assert {s.name for s in startups} == {"Acme Vector", "Beta Labs"}  # dup merged, earlier-run skipped, blank dropped
    beta = next(s for s in startups if s.name == "Beta Labs")
    assert beta.source_url is None  # URL the search never returned is not stored

    # Researched once per startup that produced someone to contact; the brief feeds the drafts.
    assert world.research_calls == len(startups)  # both startups have candidates in this run
    acme = next(s for s in startups if s.name == "Acme Vector")
    assert acme.researched_at and acme.brief["product"] == "A hosted vector search API"
    assert acme.brief["tech_signals"] == ["RAG", "vector search"]
    # A citation the model invented is kept as a fact but stripped of its made-up source.
    assert [n["url"] for n in acme.brief["recent_news"]] == ["https://news.example/acme", None]

    candidates = {c.linkedin_url: c for c in db.scalars(select(Candidate).where(Candidate.run_id == run.id))}
    # Acme: Jane + Already fill its cap of 2 (Ghost's link was never in search results).
    # Beta: Sam + Lee — Sam was never added at Acme because the cap was hit first.
    assert set(candidates) == {
        "https://www.linkedin.com/in/jane-doe",  # normalized from uk. subdomain + trailing slash
        "https://www.linkedin.com/in/already-contact",
        "https://www.linkedin.com/in/sam-founder",
        "https://www.linkedin.com/in/lee-talent",
    }
    by_startup = {s.id: s.name for s in startups}
    assert sorted(by_startup[c.startup_id] for c in candidates.values()) == ["Acme Vector", "Acme Vector", "Beta Labs", "Beta Labs"]
    assert candidates["https://www.linkedin.com/in/already-contact"].existing_contact_id is not None
    assert candidates["https://www.linkedin.com/in/jane-doe"].existing_contact_id is None

    kinds = [(e.node, e.kind) for e in _events(db, run)]
    for node, _, _ in graph.STEPS:
        assert (node, "started") in kinds and (node, "completed") in kinds, node
    assert kinds[-1] == ("agent", "completed")
    messages = " ".join(e.message for e in _events(db, run))
    assert "Skipped 1 already found in earlier runs: Old Co" in messages
    assert "ignored 1 with links not in the search results; 1 whose title names another company" in messages
    assert "https://www.linkedin.com/in/lookalike" not in candidates
    assert "Found 2 of the 3 startups requested" in messages

    db.refresh(resume)
    assert resume.parsed_profile["skills"] == ["Python", "RAG"] and resume.parsed_profile["years_experience"] == 4.0


def test_research_is_skipped_when_nobody_was_found(db, resume, world):
    """No decision-makers means no email, so the run must not pay to research the companies."""
    world.kdms = {"people": []}
    run = _make_run(db, resume)
    graph.execute_run(db, run.id)
    db.refresh(run)
    assert world.research_calls == 0
    assert not any(s.brief for s in db.scalars(select(Startup).where(Startup.run_id == run.id)))


def test_a_malformed_brief_costs_one_company_not_the_run(db, resume, world):
    # A list where a URL was asked for used to raise while checking it against the search results.
    world.brief_override = {
        "what_they_do": "Builds things.",
        "stage": True,  # a bool must not become the string "True" in a draft
        "recent_news": [{"fact": "Raised money", "url": ["https://news.example/acme"]}],
    }
    run = _make_run(db, resume)
    graph.execute_run(db, run.id)
    db.refresh(run)
    assert run.status == RunStatus.completed, run.error
    brief = next(s.brief for s in db.scalars(select(Startup).where(Startup.run_id == run.id)) if s.brief)
    assert brief["stage"] is None and brief["recent_news"] == [{"fact": "Raised money", "url": None}]


def test_person_at_two_startups_is_proposed_once(db, resume, world):
    run = _make_run(db, resume, num_startups=3, num_kdms=5)
    graph.execute_run(db, run.id)
    sam = list(db.scalars(select(Candidate).where(
        Candidate.run_id == run.id, Candidate.linkedin_url == "https://www.linkedin.com/in/sam-founder")))
    assert len(sam) == 1


def test_step_failure_marks_run_failed_with_step_name(db, resume, world):
    world.fail_on = "find_kdms"
    run = _make_run(db, resume)
    graph.execute_run(db, run.id)
    db.refresh(run)
    assert run.status == RunStatus.failed
    assert run.error == "Groq error: model overloaded"
    last = _events(db, run)[-1]
    assert (last.node, last.kind) == ("find_kdms", "failed")
    assert run.finished_at is not None


def test_unexpected_crash_is_reported_not_raised(db, resume, world, monkeypatch):
    monkeypatch.setattr(world, "complete_json", lambda *a, **k: 1 / 0)
    monkeypatch.setattr(llm, "complete_json", world.complete_json)
    run = _make_run(db, resume)
    graph.execute_run(db, run.id)
    db.refresh(run)
    assert run.status == RunStatus.failed
    assert run.error.startswith("Unexpected error in parse_resume (ZeroDivisionError)")


def test_no_startups_ends_early_without_failing(db, resume, world):
    world.startups = []
    run = _make_run(db, resume)
    graph.execute_run(db, run.id)
    db.refresh(run)
    assert run.status == RunStatus.completed
    nodes = {e.node for e in _events(db, run)}
    assert "find_kdms" not in nodes and "dedupe_against_db" not in nodes
    assert "find_kdms" not in world.llm_calls


def test_startups_from_failed_earlier_runs_can_be_picked_again(db, resume, world):
    earlier = _make_run(db, resume, status=RunStatus.failed)
    db.add(Startup(run_id=earlier.id, name="Old Co", website="https://oldco.example"))
    db.commit()
    run = _make_run(db, resume, num_startups=3)
    graph.execute_run(db, run.id)
    names = {s.name for s in db.scalars(select(Startup).where(Startup.run_id == run.id))}
    assert "Old Co" in names


def test_startup_where_nobody_was_found_can_be_picked_again(db, resume, world):
    earlier = _make_run(db, resume, status=RunStatus.completed)
    db.add(Startup(run_id=earlier.id, name="Old Co", website="https://oldco.example"))  # zero candidates
    db.commit()
    run = _make_run(db, resume, num_startups=3)
    graph.execute_run(db, run.id)
    assert "Old Co" in {s.name for s in db.scalars(select(Startup).where(Startup.run_id == run.id))}


@pytest.fixture
def stream_sessions(db, monkeypatch):
    """The stream opens its own sessions (by design); point them at the test transaction."""
    from contextlib import nullcontext

    from app import auth
    from app.routers import runs as runs_router

    monkeypatch.setattr(auth, "SessionLocal", lambda: nullcontext(db))
    monkeypatch.setattr(runs_router, "SessionLocal", lambda: nullcontext(db))


def _sse(client, url):
    kinds = []
    with client.stream("GET", url) as response:
        assert response.status_code == 200
        # Without no-transform the dashboard's proxy gzips the stream and holds events back.
        assert "no-transform" in response.headers["cache-control"]
        for line in response.iter_lines():
            if line.startswith("event: "):
                kinds.append(line[7:])
    return kinds


def test_stream_replays_events_and_ends(db, resume, world, stream_sessions, make_user, login):
    run = _make_run(db, resume)
    graph.execute_run(db, run.id)
    client = login(make_user(UserRole.viewer))
    kinds = _sse(client, f"/runs/{run.id}/stream")
    assert len(kinds) == len(_events(db, run)) + 1
    assert kinds[-2:] == ["completed", "end"]

    client.cookies.clear()
    assert client.get(f"/runs/{run.id}/stream").status_code == 401


def test_stream_for_dead_run_ends_instead_of_hanging(db, resume, stream_sessions, make_user, login):
    run = _make_run(db, resume, status=RunStatus.running)
    run.heartbeat_at = datetime.now(timezone.utc) - timedelta(minutes=30)
    db.commit()
    kinds = _sse(login(make_user(UserRole.viewer)), f"/runs/{run.id}/stream")
    assert kinds == ["failed", "end"]


def test_failed_run_leaves_no_reviewable_candidates(db, resume, world, make_user, login):
    world.fail_kdms_on_call = 2  # first startup's candidates get committed, then the second fails
    run = _make_run(db, resume)
    graph.execute_run(db, run.id)
    db.refresh(run)
    assert run.status == RunStatus.failed
    assert any("decision-maker(s) at Acme Vector" in e.message for e in _events(db, run))  # step 1 did save
    assert db.scalars(select(Candidate).where(Candidate.run_id == run.id)).first() is None

    in_progress = _make_run(db, resume, status=RunStatus.running)
    startup = Startup(run_id=in_progress.id, name="Live Co")
    db.add(startup)
    db.flush()
    db.add(Candidate(run_id=in_progress.id, startup_id=startup.id, name="Mid Run",
                     linkedin_url="https://www.linkedin.com/in/mid-run"))
    db.commit()
    listed = login(make_user(UserRole.viewer)).get("/candidates").json()
    assert all(c["run_id"] not in (str(run.id), str(in_progress.id)) for c in listed)


def test_run_stopped_elsewhere_mid_run_does_not_resurrect(db, resume, world, monkeypatch):
    run = _make_run(db, resume)
    real = world.complete_json

    def mark_failed_during_find_kdms(db_, model, prompt, temperature, max_tokens=4096, **kwargs):
        if kwargs["name"] == "find_kdms":
            # What fail_stale_runs in another process does, committed before this worker's next write.
            db.execute(Run.__table__.update().where(Run.id == run.id).values(status="failed", error="Stopped elsewhere"))
            db.commit()
        return real(db_, model, prompt, temperature, max_tokens, **kwargs)

    monkeypatch.setattr(llm, "complete_json", mark_failed_during_find_kdms)
    graph.execute_run(db, run.id)
    db.expire_all()
    run = db.get(Run, run.id)
    assert run.status == RunStatus.failed and run.error == "Stopped elsewhere"
    assert db.scalars(select(Candidate).where(Candidate.run_id == run.id)).first() is None
    assert ("agent", "completed") not in [(e.node, e.kind) for e in _events(db, run)]


def test_heartbeat_is_throttled_and_stops_a_run_that_is_no_longer_running(db, resume):
    run = _make_run(db, resume, status=RunStatus.running)
    ctx = graph.RunContext(db, run, resume)
    ctx.heartbeat()
    first = ctx.last_heartbeat
    ctx.heartbeat()  # within 30s → no new write
    assert ctx.last_heartbeat == first

    run.status = RunStatus.failed
    db.commit()
    ctx.last_heartbeat = None
    with pytest.raises(graph.RunStopped):
        ctx.heartbeat()
    with pytest.raises(graph.RunStopped):
        ctx.emit("find_kdms", "info", "should not be written")


def test_database_allows_one_active_run_per_resume(db, resume):
    from sqlalchemy.exc import IntegrityError

    _make_run(db, resume, status=RunStatus.pending)
    with pytest.raises(IntegrityError):
        _make_run(db, resume, status=RunStatus.running)
    db.rollback()


def test_run_only_executes_once(db, resume, world):
    run = _make_run(db, resume)
    graph.execute_run(db, run.id)
    calls = len(world.llm_calls)
    graph.execute_run(db, run.id)  # already completed → no-op
    assert len(world.llm_calls) == calls


def test_stale_running_run_is_failed_recent_one_kept(db, resume):
    other_resume = Resume(filename="other.pdf", storage_path="x.pdf")
    db.add(other_resume)
    db.commit()
    old = _make_run(db, resume, status=RunStatus.running)
    old.heartbeat_at = datetime.now(timezone.utc) - timedelta(minutes=30)
    fresh = _make_run(db, other_resume, status=RunStatus.running)
    fresh.heartbeat_at = datetime.now(timezone.utc)
    db.commit()
    assert graph.fail_stale_runs(db) >= 1
    db.refresh(old)
    db.refresh(fresh)
    assert old.status == RunStatus.failed and "stopped responding" in old.error
    assert fresh.status == RunStatus.running


def test_startup_fails_every_interrupted_run_even_with_fresh_heartbeat(db, resume):
    other = Resume(filename="other.pdf", storage_path="x.pdf")
    db.add(other)
    db.commit()
    running = _make_run(db, resume, status=RunStatus.running)
    running.heartbeat_at = datetime.now(timezone.utc)  # alive 0s ago, but its thread died with the process
    queued = _make_run(db, other, status=RunStatus.pending)
    db.commit()
    assert graph.fail_interrupted_runs(db) >= 2
    for run in (running, queued):
        db.refresh(run)
        assert run.status == RunStatus.failed and "restarted" in run.error


def test_submit_run_uses_daemon_workers(monkeypatch):
    import threading

    done = threading.Event()
    seen = {}

    def fake_run(run_id):
        seen.update(run_id=run_id, daemon=threading.current_thread().daemon)
        done.set()

    monkeypatch.setattr(graph, "run_in_background", fake_run)
    run_id = uuid.uuid4()
    graph.submit_run(run_id)
    assert done.wait(10)
    assert seen == {"run_id": run_id, "daemon": True}  # never blocks server shutdown


# --- API -------------------------------------------------------------------------

@pytest.fixture
def keys(db, make_user, fake_provider_test):
    from sqlalchemy import delete
    from app.models import ApiCredential

    db.execute(delete(ApiCredential))
    db.commit()
    fake_provider_test("groq")
    fake_provider_test("tavily")
    admin = make_user(UserRole.admin)
    vault.save_credential(db, "groq", {"api_key": "good-groq"}, admin)
    vault.save_credential(db, "tavily", {"api_key": "good-tavily"}, admin)


@pytest.fixture
def started(monkeypatch):
    calls = []
    monkeypatch.setattr(graph, "submit_run", lambda run_id: calls.append(run_id))
    return calls


def test_start_run_api(db, resume, keys, started, make_user, login):
    prompts.seed_defaults(db)  # starting a run needs an active prompt for every node
    client = login(make_user(UserRole.operator))
    response = client.post("/runs", json={"resume_id": str(resume.id), "num_startups": 2, "num_kdms_per_company": 3})
    assert response.status_code == 202
    body = response.json()
    assert body["status"] == "pending" and started == [uuid.UUID(body["id"])]
    run = db.get(Run, uuid.UUID(body["id"]))
    assert set(run.prompt_versions) == set(graph.PROMPT_NODES)

    again = client.post("/runs", json={"resume_id": str(resume.id)})
    assert again.status_code == 409


@pytest.mark.parametrize("payload", [{"num_startups": 0}, {"num_startups": 21}, {"num_kdms_per_company": 11}])
def test_start_run_rejects_out_of_range_params(resume, keys, started, make_user, login, payload):
    client = login(make_user(UserRole.admin))
    assert client.post("/runs", json={"resume_id": str(resume.id), **payload}).status_code == 422


def test_start_run_needs_keys_and_permission(db, resume, started, make_user, login):
    from sqlalchemy import delete
    from app.models import ApiCredential

    db.execute(delete(ApiCredential))
    db.commit()
    assert login(make_user(UserRole.viewer)).post("/runs", json={"resume_id": str(resume.id)}).status_code == 403
    response = login(make_user(UserRole.admin)).post("/runs", json={"resume_id": str(resume.id)})
    assert response.status_code == 400 and "groq is not configured" in response.json()["detail"]
    assert started == []


def test_run_detail_events_and_candidates_api(db, resume, world, make_user, login):
    run = _make_run(db, resume)
    graph.execute_run(db, run.id)
    client = login(make_user(UserRole.viewer))

    detail = client.get(f"/runs/{run.id}").json()
    assert detail["status"] == "completed" and detail["resume_filename"] == "priya.pdf"
    # No earlier run here, so Old Co is eligible; sorted by relevance 0.9 / 0.85 / 0.7.
    assert [s["name"] for s in detail["startups"]] == ["Acme Vector", "Old Co", "Beta Labs"]
    assert all(c["cv_filename"] == "priya.pdf" and c["status"] == "pending" for c in detail["candidates"])

    events = client.get(f"/runs/{run.id}/events").json()
    later = client.get(f"/runs/{run.id}/events?after={events[2]['id']}").json()
    assert len(later) == len(events) - 3

    pending = client.get(f"/candidates?status=pending&run_id={run.id}").json()
    assert len(pending) == len(detail["candidates"])
    assert client.get("/candidates?status=nonsense").status_code == 422


# --- building blocks -------------------------------------------------------------

@pytest.mark.parametrize("url, expected", [
    ("https://uk.linkedin.com/in/Thom-Wolf?trk=abc", "https://www.linkedin.com/in/thom-wolf"),
    ("linkedin.com/in/someone/details/experience/", "https://www.linkedin.com/in/someone"),
    ("https://www.linkedin.com/company/acme", None),
    ("https://linkedin.com.evil.example/in/x", None),
    ("https://notlinkedin.com/in/x", None),
    ("", None),
])
def test_linkedin_normalization(url, expected):
    assert linkedin_profile_url(url) == expected


@pytest.mark.parametrize("title, startup, website, other", [
    ("Founder at Sapling Says", "Sapling.ai", "https://sapling.ai", True),
    ("Founder at Stealth", "Acme", None, True),
    ("Engineer at Google, previously at Acme", "Acme", None, True),
    ("Co-Founder & COO at Multiverse", "Multiverse Computing", None, False),
    ("Talent Partner at Hippocratic AI | ex-Google", "Hippocratic", None, False),
    ("CTO @PowerfulMedical", "Powerful Medical", None, False),
    ("CTO at Acme Robotics", "Acme", "https://acmerobotics.com", False),
    ("Head of Talent", "Acme", None, False),
    (None, "Acme", None, False),
    ("Head of Talent at Scale", "Upscale", None, True),  # substring in the middle is not a match
    ("Engineer at Meta", "Metaview", None, True),  # prefix with a non-generic tail is not a match
    ("Recruiter at Hippocratic", "Hippocratic AI", None, False),
    ("Co-founder at Hugging-Face", "Hugging Face", None, False),  # hyphen inside the name
    ("Founder at Acme – building agents", "Acme", None, False),  # spaced dash separates
    ("Co-founder & CEO at monday.com", "Monday", "https://monday.com", False),  # web ending
    ("Founder at Notion.so", "Notion", None, False),
    ("CTO | Building AI agents at scale", "Acme", None, False),  # lowercase phrase, not a company
    ("CTO @Other", "Acme", None, True),
])
def test_title_names_other_company(title, startup, website, other):
    assert title_names_other_company(title, startup, website) is other


def test_company_and_domain_keys():
    assert company_key("Acme Vector, Inc.") == company_key("acme-vector") == "acmevector"
    assert website_domain("https://www.Acme.example/about") == "acme.example"


def test_docx_extraction(db, tmp_path):
    document = docx.Document()
    document.add_paragraph("Priya Sharma — Machine Learning Engineer " * 5)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Skills", "Python, PyTorch, RAG"
    path = tmp_path / "cv.docx"
    document.save(path)
    text, method = extract.extract_text(db, str(path))
    assert method == "docx" and "Skills | Python, PyTorch, RAG" in text


def test_scanned_pdf_falls_back_to_ocr(db, tmp_path, monkeypatch):
    pdf = pypdfium2.PdfDocument.new()
    pdf.new_page(612, 792)
    path = tmp_path / "scan.pdf"
    pdf.save(str(path))
    pdf.close()
    seen = {}

    def fake_complete(db, model, prompt, temperature, max_tokens=2048, **kwargs):
        seen.update(model=model, image=kwargs.get("image_data_url", "")[:22])
        return "Priya Sharma\nML Engineer\n" + "Built RAG pipelines. " * 20

    monkeypatch.setattr(llm, "complete", fake_complete)
    text, method = extract.extract_text(db, str(path))
    assert method == "vision_ocr" and text.startswith("Priya Sharma")
    assert seen == {"model": extract.OCR_MODEL, "image": "data:image/png;base64,"}


def test_missing_resume_file_is_a_clear_error(db):
    with pytest.raises(extract.ExtractionError, match="missing"):
        extract.extract_text(db, "storage/nope/never.pdf")


def test_llm_waits_on_rate_limit_then_succeeds(db, monkeypatch):
    monkeypatch.setattr(vault, "get_credential", lambda db, provider: {"api_key": "k"})
    sleeps, notices = [], []
    monkeypatch.setattr(llm.time, "sleep", sleeps.append)
    responses = iter([
        httpx.Response(429, headers={"x-ratelimit-reset-tokens": "1m2.5s"}, json={"error": {"message": "rate"}}),
        httpx.Response(429, headers={"retry-after": "3"}, json={"error": {"message": "rate"}}),
        httpx.Response(200, json={"choices": [{"message": {"content": '<think>hmm</think>{"ok": true}'}, "finish_reason": "stop"}],
                                  "usage": {"prompt_tokens": 10, "completion_tokens": 5}}),
    ])
    monkeypatch.setattr(llm.httpclient, "post", lambda *a, **k: next(responses))

    tel = telemetry.RunTelemetry(on_notice=notices.append)
    with telemetry.activate(tel):
        assert llm.complete_json(db, "openai/gpt-oss-120b", "x", 0, name="t") == {"ok": True}
    assert sleeps == [63.0, 3.5]
    assert len(notices) == 2 and tel.usage["llm_calls"] == 1 and tel.usage["prompt_tokens"] == 10


def test_llm_truncated_output_is_an_error(db, monkeypatch):
    monkeypatch.setattr(vault, "get_credential", lambda db, provider: {"api_key": "k"})
    monkeypatch.setattr(llm.httpclient, "post", lambda *a, **k: httpx.Response(
        200, json={"choices": [{"message": {"content": '{"a":'}, "finish_reason": "length"}], "usage": {}}))
    with pytest.raises(llm.LLMError, match="cut off"):
        llm.complete(db, "m", "x", 0, name="t")
